"""
Unit and integration tests for fund_liquidity module.
"""

import os
from copy import copy
import openpyxl
from openpyxl.styles import Font, PatternFill
import pandas as pd
import pytest

from fund_liquidity import (
    clean_security_id,
    copy_sheet_formula_as_value_with_format,
    prepare_fund_data,
    populate_fund_worksheet,
    process_fund_liquidity,
)


def test_clean_security_id():
    """Verify that clean_security_id strips single quotes and whitespace correctly."""
    s = pd.Series(["'0700", " AAPL ", "'MSFT'", 12345])
    cleaned = clean_security_id(s)
    assert cleaned.tolist() == ["0700", "AAPL", "MSFT'", "12345"]


def test_copy_sheet_formula_as_value_with_format(tmp_path):
    """Test copying worksheet as values while preserving styling, dimensions, and merged cells."""
    src_file = str(tmp_path / "src.xlsx")
    dest_file = str(tmp_path / "dest.xlsx")

    # Create source workbook
    src_wb = openpyxl.Workbook()
    ws = src_wb.active
    ws.title = "TestSheet"

    # Set cell with value and formula
    ws["A1"] = 100
    ws["A2"] = 200
    ws["A3"] = "=SUM(A1:A2)"
    ws["A3"].font = Font(name="Arial", bold=True, color="FF0000")
    ws["A3"].fill = PatternFill(start_color="00FF00", end_color="00FF00", fill_type="solid")
    ws["A3"].number_format = "#,##0.00"

    # Set merged cells
    ws["B1"] = "Merged Title"
    ws.merge_cells("B1:C1")

    # Set dimensions
    ws.column_dimensions["A"].width = 25.0
    ws.row_dimensions[3].height = 30.0

    src_wb.save(src_file)

    # Destination workbook
    dest_wb = openpyxl.Workbook()
    dest_wb.save(dest_file)

    # In openpyxl, data_only=True evaluates formulas when saved by Excel,
    # but for purely openpyxl-generated files without Excel calculation engine,
    # the cached value may be None or the assigned value.
    # We test the copy logic directly:
    result_wb = copy_sheet_formula_as_value_with_format(
        src_path_or_wb=src_file,
        sheet_name="TestSheet",
        dest_path_or_wb=dest_file,
        save_dest=True,
    )

    loaded_wb = openpyxl.load_workbook(dest_file)
    assert "TestSheet" in loaded_wb.sheetnames
    dest_ws = loaded_wb["TestSheet"]

    # Verify styling preserved
    assert dest_ws["A3"].font.bold is True
    assert dest_ws["A3"].number_format == "#,##0.00"
    assert dest_ws["A3"].fill.start_color.rgb == "0000FF00"

    # Verify dimensions
    assert dest_ws.column_dimensions["A"].width == 25.0
    assert dest_ws.row_dimensions[3].height == 30.0

    # Verify merged cells
    merged_ranges = [str(r) for r in dest_ws.merged_cells.ranges]
    assert "B1:C1" in merged_ranges


def test_prepare_fund_data():
    """Verify fund filtering, Sec Type sorting, and classification merging."""
    # Mock LiqTracking
    liq_tracking_data = {
        "Fund Name": ["Fund A", "Fund A", "Fund A", "Fund B"],
        "Security ID": ["'0700", "AAPL", "CASH", "MSFT"],
        "Security Desc": ["Tencent", "Apple", "Cash", "Microsoft"],
        "Sec Type": ["Common Stock", "Common Stock", "Cash Balance", "Common Stock"],
        "Col5": [1, 2, 3, 4],
        "Col6": [1, 2, 3, 4],
        "Col7": [1, 2, 3, 4],
        "Col8": [1, 2, 3, 4],
        "Col9": [1, 2, 3, 4],
        "Col10": [1, 2, 3, 4],
        "Col11": [1, 2, 3, 4],
        "Mkt Value L$": [1000000.0, 500000.0, 200000.0, 800000.0],
        "Col13": [1, 2, 3, 4],
        "Col14": [1, 2, 3, 4],
    }
    liq_df = pd.DataFrame(liq_tracking_data)

    # Mock Liquidity Source
    source_data = {
        "Identifier 2": ["Fund A0700", "Fund AAAPL"],
        "Liquidity Classification in the past 30 days": ["Liquid", "Liquid"],
        "Liquidity Classification in the past 5 years (5th percentile)": ["Semi-Liquid", "Liquid"],
        "Liquidity Classification in the past 5 years": ["Liquid", "Liquid"],
        "INDUSTRY_SECTOR": ["Communications", "Technology"],
    }
    source_df = pd.DataFrame(source_data)

    fund_a_data = prepare_fund_data("Fund A", liq_df, source_df)

    # Verify count and sorting
    assert len(fund_a_data) == 3
    # Common Stock should come before Cash Balance when sorted descending
    assert fund_a_data.iloc[0]["Sec Type"] == "Common Stock"
    assert fund_a_data.iloc[-1]["Sec Type"] == "Cash Balance"

    # Verify column order and rename of Industry_Sector
    assert "Industry_Sector" in fund_a_data.columns
    assert list(fund_a_data.columns)[-1] == "Industry_Sector"

    # Verify classifications merged
    row_0700 = fund_a_data[fund_a_data["Security ID"] == "'0700"].iloc[0]
    assert row_0700["Liquidity Classification in the past 30 days"] == "Liquid"
    assert row_0700["Industry_Sector"] == "Communications"

    # Cash Balance had no record in source_df, so merged fields should be NaN
    row_cash = fund_a_data[fund_a_data["Sec Type"] == "Cash Balance"].iloc[0]
    assert pd.isna(row_cash["Liquidity Classification in the past 30 days"])


def test_populate_fund_worksheet():
    """Verify imputation, formula generation, and scenario summary tables."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Standard Fund"

    columns = [
        "Fund Name", "Security ID", "Security Desc", "Sec Type",
        "C5", "C6", "C7", "C8", "C9", "C10", "C11",
        "Mkt Value L$", "C13", "C14",
        "Liquidity Classification in the past 30 days",
        "Liquidity Classification in the past 5 years (5th percentile)",
        "Liquidity Classification in the past 5 years",
        "Industry_Sector"
    ]
    rows = [
        ["Standard Fund", "0700", "Tencent", "Common Stock",
         1, 1, 1, 1, 1, 1, 1, 1000000.0, 1, 1,
         "Liquid", "Semi-Liquid", "Liquid", "Communications"],
        ["Standard Fund", "CASH", "Cash", "Cash Balance",
         1, 1, 1, 1, 1, 1, 1, 200000.0, 1, 1,
         None, None, None, "Cash"],
        ["Standard Fund", "UNKNOWN", "Missing Sec", "Common Stock",
         1, 1, 1, 1, 1, 1, 1, 300000.0, 1, 1,
         None, None, None, "Unknown"],
    ]
    fund_df = pd.DataFrame(rows, columns=columns)

    populate_fund_worksheet(
        worksheet=ws,
        fund_name="Standard Fund",
        fund_data=fund_df,
        workbook_name_ref="[Liquidity Result.xlsx]",
    )

    # 1. Check imputation
    # Row 2: Cash Balance (data row 1 in df, row 3 in Excel)
    assert ws["O3"].value == "HLI"
    assert ws["P3"].value == "HLI"
    assert ws["Q3"].value == "HLI"
    assert ws["O3"].fill.start_color.rgb is None or ws["O3"].fill.fill_type is None

    # Row 3: Missing Common Stock (data row 2 in df, row 4 in Excel)
    assert ws["O4"].value == "ILI"
    assert ws["P4"].value == "ILI"
    assert ws["Q4"].value == "ILI"
    assert ws["O4"].fill.start_color.rgb == "00FFFF00"
    assert ws["P4"].fill.start_color.rgb == "00FFFF00"
    assert ws["Q4"].fill.start_color.rgb == "00FFFF00"

    # 2. Check scenario headers
    assert ws["S1"].value == "Normal"
    assert ws["T1"].value == "Moderate"
    assert ws["U1"].value == "Stress"
    assert ws["V1"].value == "Hypothetical"

    # 3. Check formulas for row 2
    assert ws["S2"].value == "=(1-VLOOKUP(O2,'[Liquidity Result.xlsx]Hair Cut'!$A$1:$D$5,2,0))*L2"
    assert ws["T2"].value == "=(1-VLOOKUP(P2,'[Liquidity Result.xlsx]Hair Cut'!$A$1:$D$5,3,0))*L2"
    assert ws["U2"].value == "=(1-VLOOKUP(Q2,'[Liquidity Result.xlsx]Hair Cut'!$A$1:$D$5,4,0))*L2"
    assert ws["V2"].value == "=(1+IF(ISNA(VLOOKUP(R2,'Hair Cut'!$A$9:$B$10,2,0)),MIN('Hair Cut'!$B$9:$B$10),VLOOKUP(R2,'Hair Cut'!$A$9:$B$10,2,0)))*L2"

    # 4. Check summation formulas (max_row is 4, sum_row is 5)
    assert ws["S5"].value == "=SUM(S2:S4)"
    assert ws["T5"].value == "=SUM(T2:T4)"
    assert ws["U5"].value == "=SUM(U2:U4)"
    assert ws["V5"].value == "=SUM(V2:V4)"
    assert ws["L5"].value == "=SUM(L2:L4)"

    # 5. Check Scenario Summary Table at sum_row + 2 = 7
    table_row = 7
    assert ws.cell(row=table_row, column=1).value == "Standard Fund"
    assert ws.cell(row=table_row + 1, column=1).value == "Asset"
    assert ws.cell(row=table_row + 2, column=1).value == "Redemption"
    assert ws.cell(row=table_row + 3, column=1).value == "Monthly overhead"
    assert ws.cell(row=table_row + 4, column=1).value == "Liquidity Ratio"

    # Normal Scenario (col B)
    assert ws.cell(row=table_row, column=2).value == "Normal Scenario"
    assert ws.cell(row=table_row + 1, column=2).value == "=S5"
    assert ws.cell(row=table_row + 2, column=2).value == "=VLOOKUP(A2,'Redemption Summary'!A:N,14,0)"
    assert ws.cell(row=table_row + 3, column=2).value == f"=0.01*B{table_row+1}"
    assert ws.cell(row=table_row + 4, column=2).value == f"=B{table_row+1}/SUM(B{table_row+2}:B{table_row+3})"

    # Hypothetical 1 (col E)
    assert ws.cell(row=table_row, column=5).value == "Hypothetical 1"
    assert ws.cell(row=table_row + 1, column=5).value == "=V5"
    assert ws.cell(row=table_row + 2, column=5).value == "=$L$5*'Redemption Summary'!Q$2"
    assert ws.cell(row=table_row + 3, column=5).value == f"=0.01*E{table_row+1}"
    assert ws.cell(row=table_row + 4, column=5).value == f"=E{table_row+1}/SUM(E{table_row+2}:E{table_row+3})"

    # Check number formatting
    assert ws[f"B{table_row+4}"].number_format == "0.00%"
    assert ws[f"E{table_row+4}"].number_format == "0.00%"
    assert ws[f"B{table_row+1}"].number_format == "#,##0.00"


def test_mpf_fund_imputation_rule():
    """Verify that MPF funds leave non-cash missing classifications blank."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MPF Equity Fund"

    columns = [
        "Fund Name", "Security ID", "Security Desc", "Sec Type",
        "C5", "C6", "C7", "C8", "C9", "C10", "C11",
        "Mkt Value L$", "C13", "C14",
        "Liquidity Classification in the past 30 days",
        "Liquidity Classification in the past 5 years (5th percentile)",
        "Liquidity Classification in the past 5 years",
        "Industry_Sector"
    ]
    rows = [
        ["MPF Equity Fund", "CASH", "Cash", "Cash Balance",
         1, 1, 1, 1, 1, 1, 1, 200000.0, 1, 1,
         None, None, None, "Cash"],
        ["MPF Equity Fund", "0005", "HSBC", "Common Stock",
         1, 1, 1, 1, 1, 1, 1, 500000.0, 1, 1,
         None, None, None, "Financials"],
    ]
    fund_df = pd.DataFrame(rows, columns=columns)

    populate_fund_worksheet(
        worksheet=ws,
        fund_name="MPF Equity Fund",
        fund_data=fund_df,
    )

    # Cash Balance -> HLI
    assert ws["O2"].value == "HLI"
    # Non-cash for MPF -> Left blank (None), no fill
    assert ws["O3"].value is None
    assert ws["P3"].value is None
    assert ws["Q3"].value is None
    assert ws["O3"].fill.start_color.rgb is None or ws["O3"].fill.fill_type is None


def test_process_fund_liquidity_e2e(tmp_path):
    """Full end-to-end integration test creating Liquidity Result.xlsx."""
    folder = str(tmp_path)
    input_file = os.path.join(folder, "Liquidity Source.xlsx")
    redemption_file = os.path.join(folder, "Monthly Redemption.xlsx")
    haircut_file = os.path.join(folder, "Hair Cut.xlsx")
    output_file = os.path.join(folder, "Liquidity Result.xlsx")

    # 1. Create Hair Cut.xlsx
    wb_hc = openpyxl.Workbook()
    ws_hc = wb_hc.active
    ws_hc.title = "Hair Cut"
    # Haircut table A1:D5
    hc_data = [
        ["Class", "Normal", "Moderate", "Stress"],
        ["HLI", 0.05, 0.10, 0.20],
        ["MLI", 0.10, 0.20, 0.35],
        ["SLI", 0.15, 0.30, 0.50],
        ["ILI", 0.25, 0.45, 0.70],
    ]
    for row in hc_data:
        ws_hc.append(row)

    # Industry sector shocks A9:B10
    ws_hc["A9"] = "Financials"
    ws_hc["B9"] = -0.15
    ws_hc["A10"] = "Technology"
    ws_hc["B10"] = -0.25
    wb_hc.save(haircut_file)

    # 2. Create Monthly Redemption.xlsx
    wb_red = openpyxl.Workbook()
    ws_red = wb_red.active
    ws_red.title = "Redemption Summary"
    # Headers row 1: A to T
    headers_red = [f"Col{i}" for i in range(1, 21)]
    ws_red.append(headers_red)
    # Row 2: Factors for Hypothetical 1-4 at cols Q, R, S, T (cols 17, 18, 19, 20)
    row2_red = [""] * 20
    row2_red[0] = "Fund Alpha"
    row2_red[9] = 100000.0   # Col J (10) - Severe
    row2_red[10] = 50000.0   # Col K (11) - Moderate
    row2_red[13] = 20000.0   # Col N (14) - Normal
    row2_red[16] = 0.05      # Col Q (17) - Hypo 1
    row2_red[17] = 0.10      # Col R (18) - Hypo 2
    row2_red[18] = 0.15      # Col S (19) - Hypo 3
    row2_red[19] = 0.20      # Col T (20) - Hypo 4
    ws_red.append(row2_red)
    wb_red.save(redemption_file)

    # 3. Create Liquidity Source.xlsx
    wb_src = openpyxl.Workbook()

    # Sheet: Equity Fund List
    ws_fund_list = wb_src.active
    ws_fund_list.title = "Equity Fund List"
    ws_fund_list.append(["Fund Name"])
    ws_fund_list.append(["Fund Alpha"])
    ws_fund_list.append(["MPF Fund Beta"])

    # Sheet: LiqTracking
    ws_liq_track = wb_src.create_sheet(title="LiqTracking")
    track_headers = [
        "Fund Name", "Security ID", "Security Desc", "Sec Type",
        "Q1", "Q2", "Q3", "Q4", "Q5", "Q6", "Q7",
        "Mkt Value L$", "Extra1", "Extra2"
    ]
    ws_liq_track.append(track_headers)
    ws_liq_track.append(["Fund Alpha", "'0700", "Tencent", "Common Stock", 1, 1, 1, 1, 1, 1, 1, 1000000.0, 1, 1])
    ws_liq_track.append(["Fund Alpha", "'CASH", "Cash Bal", "Cash Balance", 1, 1, 1, 1, 1, 1, 1, 200000.0, 1, 1])
    ws_liq_track.append(["MPF Fund Beta", "'0005", "HSBC", "Common Stock", 1, 1, 1, 1, 1, 1, 1, 800000.0, 1, 1])

    # Sheet: Liquidity Source
    ws_liq_src = wb_src.create_sheet(title="Liquidity Source")
    liq_src_headers = [
        "Fund Name", "Security ID", "Identifier 2",
        "Liquidity Classification in the past 30 days",
        "Liquidity Classification in the past 5 years (5th percentile)",
        "Liquidity Classification in the past 5 years",
        "INDUSTRY_SECTOR"
    ]
    ws_liq_src.append(liq_src_headers)
    ws_liq_src.append(["Fund Alpha", "0700", "Fund Alpha0700", "Liquid", "Semi-Liquid", "Liquid", "Technology"])
    wb_src.save(input_file)

    # Run the processor
    result_path = process_fund_liquidity(
        folder=folder,
        input_file=input_file,
        output_file=output_file,
        redemption_file=redemption_file,
        haircut_file=haircut_file,
    )

    assert os.path.exists(result_path)
    wb_res = openpyxl.load_workbook(result_path)

    # Check sheets exist
    assert "Redemption Summary" in wb_res.sheetnames
    assert "Hair Cut" in wb_res.sheetnames
    assert "Fund Alpha" in wb_res.sheetnames
    assert "MPF Fund Beta" in wb_res.sheetnames

    # Check Fund Alpha tab
    ws_alpha = wb_res["Fund Alpha"]
    assert ws_alpha["S1"].value == "Normal"
    # Row 2 (Tencent) has merged classifications
    assert ws_alpha["O2"].value == "Liquid"
    # Row 3 (Cash) has imputed HLI
    assert ws_alpha["O3"].value == "HLI"

    # Check MPF Fund Beta tab
    ws_mpf = wb_res["MPF Fund Beta"]
    # Row 2 (HSBC) is non-cash MPF -> left blank
    assert ws_mpf["O2"].value is None

    # Check Scenario Summary Table in Fund Alpha
    # sum_row is 4 (2 data rows + header = max_row 3, sum_row 4)
    # table_row is 6
    assert ws_alpha["A6"].value == "Fund Alpha"
    assert ws_alpha["B6"].value == "Normal Scenario"
    assert ws_alpha["B7"].value == "=S4"


def test_missing_input_file_raises():
    """Verify that missing files raise FileNotFoundError."""
    with pytest.raises(FileNotFoundError):
        process_fund_liquidity(
            input_file="non_existent_input.xlsx",
            redemption_file="non_existent_red.xlsx",
            haircut_file="non_existent_hc.xlsx",
        )


def test_empty_fund_handling(tmp_path):
    """Verify handling when a fund in Equity Fund List has no records in LiqTracking."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Empty Fund"

    empty_df = pd.DataFrame(columns=[
        "Fund Name", "Security ID", "Security Desc", "Sec Type",
        "Mkt Value L$", "Liquidity Classification in the past 30 days",
        "Liquidity Classification in the past 5 years (5th percentile)",
        "Liquidity Classification in the past 5 years", "Industry_Sector"
    ])

    populate_fund_worksheet(
        worksheet=ws,
        fund_name="Empty Fund",
        fund_data=empty_df,
    )

    # Headers present
    assert ws["A1"].value == "Fund Name"
    assert ws["S1"].value == "Normal"
    # sum_row is 2
    assert ws["S2"].value == 0
    # Scenario table at row 4
    assert ws["A4"].value == "Empty Fund"
    assert ws["B4"].value == "Normal Scenario"


def test_use_internal_ref_formulas(tmp_path):
    """Verify formulas generated with use_internal_ref=True have no workbook prefix."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test Internal Ref"

    df = pd.DataFrame({
        "Fund Name": ["Fund X"],
        "Security ID": ["AAPL"],
        "Security Desc": ["Apple"],
        "Sec Type": ["Common Stock"],
        "C5": [1], "C6": [1], "C7": [1], "C8": [1], "C9": [1], "C10": [1], "C11": [1],
        "Mkt Value L$": [10000.0], "C13": [1], "C14": [1],
        "Liquidity Classification in the past 30 days": ["Liquid"],
        "Liquidity Classification in the past 5 years (5th percentile)": ["Liquid"],
        "Liquidity Classification in the past 5 years": ["Liquid"],
        "Industry_Sector": ["Technology"],
    })

    populate_fund_worksheet(
        worksheet=ws,
        fund_name="Fund X",
        fund_data=df,
        workbook_name_ref="",
        format_numbers=False,
    )

    # Check formula in S2 has 'Hair Cut' without [Liquidity Result.xlsx]
    assert ws["S2"].value == "=(1-VLOOKUP(O2,'Hair Cut'!$A$1:$D$5,2,0))*L2"
    assert ws["T2"].value == "=(1-VLOOKUP(P2,'Hair Cut'!$A$1:$D$5,3,0))*L2"
    assert ws["U2"].value == "=(1-VLOOKUP(Q2,'Hair Cut'!$A$1:$D$5,4,0))*L2"
    # Without format_numbers, General formatting is retained
    assert ws["B6"].number_format == "General"


def test_cli_parsing():
    """Verify CLI argument parsing."""
    from fund_liquidity import parse_args
    import sys

    test_args = [
        "fund_liquidity.py",
        "-i", "custom_in.xlsx",
        "-o", "custom_out.xlsx",
        "-r", "custom_red.xlsx",
        "-c", "custom_hc.xlsx",
        "--use-internal-ref",
        "--no-format-numbers",
    ]
    orig_argv = sys.argv
    try:
        sys.argv = test_args
        args = parse_args()
        assert args.input_file == "custom_in.xlsx"
        assert args.output_file == "custom_out.xlsx"
        assert args.monthly_redemption == "custom_red.xlsx"
        assert args.hair_cut == "custom_hc.xlsx"
        assert args.use_internal_ref is True
        assert args.no_format_numbers is True
    finally:
        sys.argv = orig_argv

