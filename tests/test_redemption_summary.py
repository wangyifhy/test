"""Unit and integration tests for redemption_summary."""

from __future__ import annotations

from datetime import datetime

import pytest
from openpyxl import Workbook, load_workbook

from redemption_summary import (
    add_months,
    build_month_index,
    month_label,
    parse_args,
    parse_month,
    parse_number,
    percentile_inc,
    percentile_month,
    process_redemption_summary,
    summarize_fund,
    write_summary,
)


def test_add_months_and_label():
    assert add_months(datetime(2026, 1, 1), -1) == datetime(2025, 12, 1)
    assert add_months(datetime(2025, 12, 1), 1) == datetime(2026, 1, 1)
    assert month_label(datetime(2026, 8, 15)) == "2026 Aug"


@pytest.mark.parametrize(
    "raw, expected",
    [
        (datetime(2026, 8, 31, 15, 0), datetime(2026, 8, 1)),
        ("2026-08-31", datetime(2026, 8, 1)),
        ("2026-08", datetime(2026, 8, 1)),
        ("2026 Aug", datetime(2026, 8, 1)),
        ("2026 August", datetime(2026, 8, 1)),
    ],
)
def test_parse_month(raw, expected):
    assert parse_month(raw) == expected


def test_parse_month_rejects_empty():
    with pytest.raises(ValueError):
        parse_month(None)
    with pytest.raises(ValueError):
        parse_month("not-a-date")


@pytest.mark.parametrize(
    "raw, expected",
    [
        (None, None),
        ("", None),
        ("-", None),
        ("—", None),
        ("1,234.50", 1234.50),
        (100, 100.0),
        ("  12.5 ", 12.5),
    ],
)
def test_parse_number(raw, expected):
    assert parse_number(raw) == expected


def test_percentile_inc_matches_excel():
    # PERCENTILE.INC([1,2,3,4,5], 0.95) = 4.8
    assert percentile_inc([1, 2, 3, 4, 5], 0.95) == pytest.approx(4.8)
    assert percentile_inc([10], 0.95) == 10.0
    assert percentile_inc([], 0.95) == 0.0
    # 36-point window: k = 1 + 35*0.95 = 34.25 -> interpolate 34th and 35th
    values = list(range(1, 37))
    assert percentile_inc(values, 0.95) == pytest.approx(34.25)


def test_percentile_month_uses_upper_order_statistic():
    series = [(f"m{i}", float(i)) for i in range(1, 6)]
    # k = 1 + 4*0.95 = 4.8, ceil -> rank 5
    assert percentile_month(series, 0.95) == "m5"


def test_percentile_month_stable_on_ties():
    series = [("2024 Jan", 10.0), ("2024 Feb", 20.0), ("2024 Mar", 20.0)]
    # k = 1 + 2*0.95 = 2.9, ceil -> rank 3; stable sort keeps Feb before Mar at 20
    assert percentile_month(series, 0.95) == "2024 Mar"


def test_summarize_fund_missing_months_are_zero():
    labels = [month_label(add_months(datetime(2026, 8, 1), -i)) for i in range(35, -1, -1)]
    by_fund = {
        "Fund A": {
            "2026 Aug": 100.0,
            "2026 Jul": 40.0,
            "2025 Aug": 80.0,
            "2024 Jan": 200.0,
        }
    }
    stats = summarize_fund("Fund A", labels, by_fund)
    assert stats["max"] == 200.0
    assert stats["max_month"] == "2024 Jan"
    assert stats["current"] == 100.0
    # Last 12 months ending Aug 2026 are Sep 2025..Aug 2026 (Aug 2025 is excluded)
    assert stats["avg_1y"] == pytest.approx(140.0 / 12)
    # Unknown fund is all zeros
    missing = summarize_fund("Ghost", labels, by_fund)
    assert missing["max"] == 0.0
    assert missing["current"] == 0.0
    assert missing["avg_1y"] == 0.0


def test_max_month_keeps_earliest_tie():
    labels = ["2024 Jan", "2024 Feb", "2024 Mar"]
    by_fund = {"F": {"2024 Jan": 5.0, "2024 Feb": 9.0, "2024 Mar": 9.0}}
    stats = summarize_fund("F", labels, by_fund)
    assert stats["max"] == 9.0
    assert stats["max_month"] == "2024 Feb"


def _write_month_sheet(wb, label, rows):
    ws = wb.create_sheet(label)
    ws["A1"] = "Portfolio"
    ws["E1"] = "Outflow Base$"
    for i, (fund, amount) in enumerate(rows, start=2):
        ws.cell(i, 1).value = fund
        ws.cell(i, 5).value = amount


def _make_input_file(path, current, funds):
    wb = Workbook()
    ws = wb.active
    ws["B1"] = current
    ws["D1"] = "Fund List"
    ws["F1"] = "Base Currency"
    for i, (fund, ccy) in enumerate(funds, start=2):
        ws.cell(i, 4).value = fund
        ws.cell(i, 6).value = ccy
    wb.save(path)


def _make_redemption_file(path, months, extra_old_fund=True):
    wb = Workbook()
    summary = wb.active
    summary.title = "Redemption Summary"
    summary["A1"] = "Summary"
    summary["Q1"] = "Hypo 1"
    summary["R1"] = 0.10
    summary["S1"] = "Hypo 2"
    summary["T1"] = 0.20
    summary["Q9"] = "FX"
    summary["R10"] = 7.8  # HKD
    summary["R11"] = 7.2  # CNY
    if extra_old_fund:
        summary["A2"] = "OLD FUND"
        summary["B2"] = "2020 Jan"
        summary["C2"] = 999
        summary["I2"] = "=B2"
        summary["A5"] = "STALE"
        summary["I5"] = "=B5"
    for label, rows in months.items():
        _write_month_sheet(wb, label, rows)
    wb.save(path)


def test_end_to_end_updates_summary_and_preserves_fx(tmp_path):
    input_path = tmp_path / "Input File.xlsx"
    redemption_path = tmp_path / "Monthly Redemption.xlsx"

    current = datetime(2026, 8, 1)
    _make_input_file(
        input_path,
        current,
        [("Alpha", "USD"), ("Beta", "CNY"), ("Gamma", "HKD")],
    )

    months = {}
    # Sparse history: only a few tabs exist; missing months / funds -> 0
    months["2026 Aug"] = [("Alpha", 50), ("Beta", "1,200"), ("Gamma", "-")]
    months["2026 Jul"] = [("Alpha", 80), ("Beta", 400)]
    months["2025 Aug"] = [("Alpha", 20), ("Other", 9999)]
    months["2024 Jan"] = [("Alpha", 300), ("Beta", 100)]
    months["2023 Sep"] = [("Alpha", 10)]  # first month of the 36-month window
    months["2023 Aug"] = [("Alpha", 99999)]  # one month before the window; ignored
    _make_redemption_file(redemption_path, months)

    out_path, rows = process_redemption_summary(
        input_file=str(input_path),
        redemption_file=str(redemption_path),
    )
    assert out_path == str(redemption_path)
    assert [r["fund"] for r in rows] == ["Alpha", "Beta", "Gamma"]

    alpha = next(r for r in rows if r["fund"] == "Alpha")
    assert alpha["max"] == 300.0
    assert alpha["max_month"] == "2024 Jan"
    assert alpha["current"] == 50.0
    # last 12 months (Sep 2025..Aug 2026): Jul=80, Aug=50; Aug 2025 is excluded
    assert alpha["avg_1y"] == pytest.approx(130.0 / 12)

    beta = next(r for r in rows if r["fund"] == "Beta")
    assert beta["current"] == 1200.0
    assert beta["max"] == 1200.0
    assert beta["currency"] == "CNY"

    gamma = next(r for r in rows if r["fund"] == "Gamma")
    assert gamma["current"] == 0.0  # dash treated as missing -> 0
    assert gamma["max"] == 0.0

    wb = load_workbook(redemption_path)
    ws = wb["Redemption Summary"]
    # Old fund rows cleared
    assert ws["A2"].value == "Alpha"
    assert ws["A5"].value is None
    assert ws["I5"].value is None
    # Layout: B = max month, C = max value (headers historically swapped)
    assert ws["B2"].value == "2024 Jan"
    assert ws["C2"].value == 300.0
    assert ws["H3"].value == "CNY"
    assert ws["I2"].value == "=B2"
    assert ws["J2"].value == '=IF($H2<>"USD",IF($H2="CNY",C2/$R$11,C2/$R$10),C2)'
    assert ws["K2"].value == '=IF($H2<>"USD",IF($H2="CNY",D2/$R$11,D2/$R$10),D2)'
    assert ws["L2"].value == "=E2"
    assert ws["M2"].value == '=IF($H2<>"USD",IF($H2="CNY",F2/$R$11,F2/$R$10),F2)'
    assert ws["N2"].value == '=IF($H2<>"USD",IF($H2="CNY",G2/$R$11,G2/$R$10),G2)'
    # Hypo / FX cells preserved
    assert ws["Q1"].value == "Hypo 1"
    assert ws["R1"].value == 0.10
    assert ws["T1"].value == 0.20
    assert ws["R10"].value == 7.8
    assert ws["R11"].value == 7.2

    labels, by_fund = build_month_index(wb, current)
    assert labels[0] == "2023 Sep"
    assert labels[-1] == "2026 Aug"
    assert len(labels) == 36
    # Sheet exists but is outside? Wait, 2023 Sep is exactly the first month of the window
    # (Aug 2026 minus 35 months = Sep 2023). So it IS in the window.
    assert "2023 Sep" in by_fund.get("Alpha", {})


def test_write_summary_does_not_overwrite_existing_headers(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Custom"
    ws["B1"] = "Keep"
    write_summary(
        ws,
        [
            {
                "fund": "X",
                "max_month": "2026 Jan",
                "max": 1,
                "p95": 1,
                "p95_month": "2026 Jan",
                "current": 1,
                "avg_1y": 1,
                "currency": "USD",
            }
        ],
    )
    assert ws["A1"].value == "Custom"
    assert ws["B1"].value == "Keep"
    assert ws["A2"].value == "X"


def test_missing_files_raise(tmp_path):
    with pytest.raises(FileNotFoundError):
        process_redemption_summary(
            input_file=str(tmp_path / "missing.xlsx"),
            redemption_file=str(tmp_path / "also-missing.xlsx"),
        )


def test_parse_args_paths():
    args = parse_args(
        ["--folder", "/tmp/liq", "--input-file", "in.xlsx", "--redemption-file", "out.xlsx"]
    )
    assert args.folder == "/tmp/liq"
    assert args.input_file == "in.xlsx"
    assert args.redemption_file == "out.xlsx"


def test_missing_summary_sheet_raises(tmp_path):
    input_path = tmp_path / "Input File.xlsx"
    redemption_path = tmp_path / "Monthly Redemption.xlsx"
    _make_input_file(input_path, datetime(2026, 8, 1), [("A", "USD")])
    wb = Workbook()
    wb.active.title = "Not Summary"
    wb.save(redemption_path)
    with pytest.raises(ValueError, match="Redemption Summary"):
        process_redemption_summary(
            input_file=str(input_path),
            redemption_file=str(redemption_path),
        )
