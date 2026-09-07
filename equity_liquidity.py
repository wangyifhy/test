#!/usr/bin/env python3
"""
Equity Liquidity Calculation and Sheet Updater

Updates the 'Liquidity' tab in the Master Equity Liquidity workbook based on:
1. 'BBG Formula Save as Value' (Bloomberg historical data)
2. 'LiqTracking' (Market USD position tracking)
3. 'Summary' (VLOOKUP classification criteria)

Key Optimizations & Improvements:
- Resolves Bloomberg error string issues ('#N/A N/A') via robust numeric coercion.
- Prevents row explosion and length mismatch when LiqTracking contains duplicate identifiers.
- Single-pass Excel write and openpyxl cell formatting, eliminating triple file re-saves.
- Vectorized Excel formula generation for fast processing.
- Configurable CLI interface with logging and error handling.
"""

import argparse
import logging
import os
import sys
from typing import Optional

import numpy as np
import pandas as pd
from openpyxl import load_workbook

# Default network/local path
DEFAULT_FILE_PATH = r"Q:\Risk Management\Liquidity\Master File\Equity\Equity Liquidity.xlsx"

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    datefmt="%Y-%m-%d %H:%M:%S",
)
logger = logging.getLogger(__name__)


def read_bbg_data(file_path: str, sheet_name: str = "BBG Formula Save as Value") -> pd.DataFrame:
    """Read the Bloomberg value export sheet."""
    logger.info("Reading sheet '%s' from '%s'...", sheet_name, file_path)
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Workbook not found at path: {file_path}")

    df = pd.read_excel(file_path, sheet_name=sheet_name, header=0)
    logger.info("Loaded '%s' with %d rows and %d columns.", sheet_name, len(df), len(df.columns))
    return df


def read_liq_tracking_data(file_path: str, sheet_name: str = "LiqTracking") -> pd.DataFrame:
    """Read market USD data from the LiqTracking sheet."""
    logger.info("Reading sheet '%s' from '%s'...", sheet_name, file_path)
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Workbook not found at path: {file_path}")

    tracking_df = (
        pd.read_excel(
            file_path,
            sheet_name=sheet_name,
            skiprows=6,
            header=None,
            usecols=[0, 1, 11],
        )
        .assign(
            Identifier1=lambda x: x[0].astype(str) + x[1].astype(str).str.lstrip("'")
        )[[ "Identifier1", 11 ]]
        .rename(columns={11: "mkt_USD"})
    )

    # Coerce mkt_USD to numeric
    tracking_df["mkt_USD"] = pd.to_numeric(tracking_df["mkt_USD"], errors="coerce")

    # If duplicate identifiers exist, aggregate by sum to avoid merge row explosion
    if tracking_df["Identifier1"].duplicated().any():
        dup_count = tracking_df["Identifier1"].duplicated().sum()
        logger.warning(
            "Found %d duplicate Identifier1 rows in '%s'. Aggregating mkt_USD by sum.",
            dup_count,
            sheet_name,
        )
        tracking_df = (
            tracking_df.groupby("Identifier1", as_index=False)["mkt_USD"]
            .sum(min_count=1)
        )

    logger.info(
        "Loaded '%s' tracking records: %d unique identifiers.",
        sheet_name,
        len(tracking_df),
    )
    return tracking_df


def build_liquidity_dataframe(
    bbg_df: pd.DataFrame,
    tracking_df: pd.DataFrame,
    use_iferror: bool = False,
) -> pd.DataFrame:
    """
    Construct the consolidated Liquidity DataFrame with calculated trading volumes,
    market USD values, and Excel formulas for percentage and classification lookups.
    """
    logger.info("Building liquidity dataframe...")
    if len(bbg_df.columns) < 7:
        raise ValueError(
            f"BBG sheet requires at least 7 columns, but only found {len(bbg_df.columns)}."
        )

    # 1. Copy Columns A to F (first 6 columns)
    liquidity = bbg_df.iloc[:, :6].copy()

    # Determine date columns (Cols G to BN, corresponding to indices 6 to 66)
    max_date_col = min(len(bbg_df.columns), 66)
    date_columns = bbg_df.columns[6:max_date_col]
    logger.info("Analyzing %d volume date columns (columns 6 to %d).", len(date_columns), max_date_col)

    # Ensure numeric conversion (handles Bloomberg errors like '#N/A N/A')
    numeric_date_data = bbg_df[date_columns].apply(pd.to_numeric, errors="coerce")

    # 2. Extract Trading Volume metrics
    # Column G (7th column, index 6) is Trading Volume in the past 30 Days
    trading_volume_30d = pd.to_numeric(bbg_df.iloc[:, 6], errors="coerce").fillna(0)
    least_trading_volume = numeric_date_data.min(axis=1).fillna(0)
    fifth_percentile_volume = numeric_date_data.quantile(0.05, axis=1).fillna(0)

    # 3. Industry Sector & Exchange Code
    if "INDUSTRY_SECTOR" in bbg_df.columns:
        industry = bbg_df[["INDUSTRY_SECTOR"]].copy()
    else:
        logger.warning("Column 'INDUSTRY_SECTOR' not found; filling with NaN.")
        industry = pd.DataFrame({"INDUSTRY_SECTOR": [np.nan] * len(bbg_df)})

    if "EXCH_CODE" in bbg_df.columns:
        exch_code = bbg_df["EXCH_CODE"].values
    else:
        logger.warning("Column 'EXCH_CODE' not found; filling with NaN.")
        exch_code = [np.nan] * len(bbg_df)

    # 4. Map mkt_USD using Identifier1
    liquidity_identifier = (
        bbg_df.iloc[:, 0].astype(str) + bbg_df.iloc[:, 1].astype(str)
    )
    tracking_map = tracking_df.set_index("Identifier1")["mkt_USD"].to_dict()
    mapped_mkt_usd = liquidity_identifier.map(tracking_map)

    # 5. Build combined columns
    n_rows = len(bbg_df)
    row_indices = np.arange(2, n_rows + 2)  # Excel 1-based indexing, header is row 1

    if use_iferror:
        pct_30d = [f'=IFERROR(F{i}/G{i}, "")' for i in row_indices]
        pct_5y = [f'=IFERROR(F{i}/J{i}, "")' for i in row_indices]
        pct_5p = [f'=IFERROR(F{i}/M{i}, "")' for i in row_indices]
    else:
        pct_30d = [f"=F{i}/G{i}" for i in row_indices]
        pct_5y = [f"=F{i}/J{i}" for i in row_indices]
        pct_5p = [f"=F{i}/M{i}" for i in row_indices]

    vlookup_30d = [f"=VLOOKUP(H{i},Summary!$I$2:$K$5,3,1)" for i in row_indices]
    vlookup_5y = [f"=VLOOKUP(K{i},Summary!$I$2:$K$5,3,1)" for i in row_indices]
    vlookup_5p = [f"=VLOOKUP(N{i},Summary!$I$2:$K$5,3,1)" for i in row_indices]

    new_metrics = pd.DataFrame(
        {
            "Trading Volume in the past 30 Days": trading_volume_30d,
            "% (Past 30 Days)": pct_30d,
            "Liquidity Classification in the past 30 days": vlookup_30d,
            "Least Trading Volume in the past 5 years": least_trading_volume,
            "% (Past 5 Years)": pct_5y,
            "Liquidity Classification in the past 5 years": vlookup_5y,
            "5th percentile Volume in the past 5 years": fifth_percentile_volume,
            "% (5th percentile in the past 5 years)": pct_5p,
            "Liquidity Classification in the past 5 years (5th percentile)": vlookup_5p,
        }
    )

    combined_liquidity = pd.concat([liquidity, new_metrics, industry], axis=1)
    combined_liquidity["mkt_USD"] = mapped_mkt_usd.values
    combined_liquidity["EXCH_CODE"] = exch_code

    logger.info("Combined liquidity dataframe constructed: %d rows, %d columns.", len(combined_liquidity), len(combined_liquidity.columns))
    return combined_liquidity


def update_liquidity_sheet(
    file_path: str,
    combined_liquidity: pd.DataFrame,
    sheet_name: str = "Liquidity",
) -> None:
    """
    Write the combined liquidity DataFrame to the target workbook sheet in a single pass,
    applying percentage formatting to columns H (8), K (11), and N (14).
    """
    logger.info("Writing updated '%s' sheet to '%s'...", sheet_name, file_path)
    mode = "a" if os.path.exists(file_path) else "w"
    writer_kwargs = {"engine": "openpyxl"}
    if mode == "a":
        writer_kwargs["mode"] = "a"
        writer_kwargs["if_sheet_exists"] = "replace"

    with pd.ExcelWriter(file_path, **writer_kwargs) as writer:
        combined_liquidity.to_excel(writer, sheet_name=sheet_name, index=False)

        # Access the worksheet to format percentage columns
        worksheet = writer.sheets[sheet_name]
        # Columns H, K, N correspond to 1-indexed column indices 8, 11, 14
        pct_columns = [8, 11, 14]
        for col_idx in pct_columns:
            for row in worksheet.iter_rows(
                min_row=2,
                max_row=worksheet.max_row,
                min_col=col_idx,
                max_col=col_idx,
            ):
                for cell in row:
                    cell.number_format = "0.00%"

    logger.info("Successfully updated sheet '%s' with percentage formatting.", sheet_name)


def process_equity_liquidity(
    file_path: str,
    bbg_sheet: str = "BBG Formula Save as Value",
    tracking_sheet: str = "LiqTracking",
    output_sheet: str = "Liquidity",
    use_iferror: bool = False,
) -> pd.DataFrame:
    """Execute the full liquidity update pipeline."""
    logger.info("Starting Equity Liquidity update process...")
    bbg_df = read_bbg_data(file_path, sheet_name=bbg_sheet)
    tracking_df = read_liq_tracking_data(file_path, sheet_name=tracking_sheet)
    combined_df = build_liquidity_dataframe(bbg_df, tracking_df, use_iferror=use_iferror)
    update_liquidity_sheet(file_path, combined_df, sheet_name=output_sheet)
    logger.info("The '%s' sheet has been successfully updated.", output_sheet)
    return combined_df


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Update the 'Liquidity' sheet in the Equity Liquidity Master Workbook."
    )
    parser.add_argument(
        "-f",
        "--file-path",
        default=DEFAULT_FILE_PATH,
        help=f"Path to Equity Liquidity Excel file (default: {DEFAULT_FILE_PATH})",
    )
    parser.add_argument(
        "--bbg-sheet",
        default="BBG Formula Save as Value",
        help="Source sheet name for Bloomberg formula values (default: 'BBG Formula Save as Value')",
    )
    parser.add_argument(
        "--tracking-sheet",
        default="LiqTracking",
        help="Source sheet name for liquidity tracking / mkt_USD (default: 'LiqTracking')",
    )
    parser.add_argument(
        "--output-sheet",
        default="Liquidity",
        help="Target sheet name to update (default: 'Liquidity')",
    )
    parser.add_argument(
        "--use-iferror",
        action="store_true",
        help="Wrap Excel division formulas with IFERROR to guard against division by zero",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    try:
        process_equity_liquidity(
            file_path=args.file_path,
            bbg_sheet=args.bbg_sheet,
            tracking_sheet=args.tracking_sheet,
            output_sheet=args.output_sheet,
            use_iferror=args.use_iferror,
        )
    except Exception as exc:
        logger.error("Execution failed: %s", exc, exc_info=True)
        sys.exit(1)


if __name__ == "__main__":
    main()
